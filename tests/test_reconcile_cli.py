"""Pruebas del uso desde la terminal y del Excel."""

from pathlib import Path

from openpyxl import load_workbook

from reconcile import main

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "sample_data"
WITH_ERROR = ROOT / "sample_data_error"


def test_cli_prints_reconciliation_and_writes_excel(tmp_path, capsys):
    output = tmp_path / "report.xlsx"
    code = main([str(SAMPLES / "bank.csv"), str(SAMPLES / "books.csv"), "--output", str(output)])
    assert code == 0

    printed = capsys.readouterr().out
    assert "balances.json" in printed          # tomó los saldos de los datos de ejemplo
    assert "= Adjusted bank balance" in printed and "= Adjusted book balance" in printed
    assert "RECONCILED: the difference is $0.00." in printed

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Summary", "Matched", "Review", "Unmatched bank", "Unmatched books"]
    assert workbook["Matched"]["A1"].value == "Status"
    assert workbook["Matched"]["E2"].number_format.startswith('"$"')
    labels = [row[0] for row in workbook["Summary"].iter_rows(values_only=True)]
    assert "+ Deposits in transit" in labels and "Difference" in labels
    # 3 movimientos solo en el banco y 7 asientos solo en los libros (+ encabezado)
    assert workbook["Unmatched bank"].max_row == 4
    assert workbook["Unmatched books"].max_row == 8


def test_cli_with_error_sample_explains_difference(tmp_path, capsys):
    code = main([str(WITH_ERROR / "bank.csv"), str(WITH_ERROR / "books.csv"), "-o", str(tmp_path / "r.xlsx")])
    assert code == 0
    printed = capsys.readouterr().out
    assert "NOT RECONCILED" in printed and "-$90.00" in printed and "swapped" in printed


def test_cli_balances_given_by_hand_override_the_file(tmp_path, capsys):
    code = main([str(SAMPLES / "bank.csv"), str(SAMPLES / "books.csv"),
                 "--bank-balance", "$93,505.90", "--book-balance", "101,962.94", "-o", str(tmp_path / "r.xlsx")])
    assert code == 0
    printed = capsys.readouterr().out
    assert "balances.json" not in printed and "RECONCILED" in printed


def test_cli_without_balances(tmp_path, capsys):
    # Copiamos los archivos a una carpeta sin balances.json
    (tmp_path / "bank.csv").write_text((SAMPLES / "bank.csv").read_text())
    (tmp_path / "books.csv").write_text((SAMPLES / "books.csv").read_text())
    code = main([str(tmp_path / "bank.csv"), str(tmp_path / "books.csv"), "-o", str(tmp_path / "r.xlsx")])
    assert code == 0
    assert "Give both ending balances" in capsys.readouterr().out


def test_cli_works_with_debit_credit_bank_file(tmp_path, capsys):
    code = main([str(SAMPLES / "bank_debit_credit.csv"), str(SAMPLES / "books.csv"),
                 "--output", str(tmp_path / "r.xlsx")])
    assert code == 0
    assert "RECONCILED: the difference is $0.00." in capsys.readouterr().out


def test_cli_reports_bad_file_without_crashing(tmp_path, capsys):
    bad = tmp_path / "bad.csv"
    bad.write_text("Something,Else\n1,2\n")
    code = main([str(bad), str(SAMPLES / "books.csv"), "--output", str(tmp_path / "r.xlsx")])
    assert code == 1
    assert "date column" in capsys.readouterr().err


def test_cli_bad_balance(tmp_path, capsys):
    code = main([str(SAMPLES / "bank.csv"), str(SAMPLES / "books.csv"), "--bank-balance", "lots"])
    assert code == 1
    assert "Error" in capsys.readouterr().err


def test_cli_missing_file(tmp_path, capsys):
    code = main([str(tmp_path / "nope.csv"), str(SAMPLES / "books.csv")])
    assert code == 1
    assert "Error" in capsys.readouterr().err
