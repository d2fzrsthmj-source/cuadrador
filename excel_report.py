"""Exporta el resultado de la conciliación a un archivo Excel.

Hojas: Reconciliation (formato estándar, lista para imprimir), Summary, Matched,
Review, Unmatched bank, Unmatched books.
Los textos van en inglés porque el Excel es para los clientes (negocios de EE. UU.).
"""

from itertools import zip_longest

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Formato de dólares: negativos en rojo con signo menos
MONEY_FORMAT = '"$"#,##0.00;[Red]-"$"#,##0.00'
DATE_FORMAT = "mm/dd/yyyy"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
MAX_COLUMN_WIDTH = 70
THIN_TOP = Border(top=Side(style="thin"))
DOUBLE_BOTTOM = Border(top=Side(style="thin"), bottom=Side(style="double"))
GOOD_FONT = Font(bold=True, color="166534")
BAD_FONT = Font(bold=True, color="991B1B")

MATCH_HEADERS = ["Status", "Bank line", "Bank date", "Bank description", "Bank amount",
                 "Book ref", "Book date", "Book name", "Book amount", "Reason"]


def match_rows(match):
    """Una fila por cada pareja banco/libros del grupo (un depósito de 2 asientos ocupa 2 filas)."""
    rows = []
    for bank, entry in zip_longest(match.bank, match.books):
        rows.append([
            match.status,
            bank.line if bank else None,
            bank.date if bank else None,
            bank.description if bank else None,
            bank.amount if bank else None,
            entry.reference if entry else None,
            entry.date if entry else None,
            entry.description if entry else None,
            entry.amount if entry else None,
            match.reason,
        ])
    return rows


def style_header(row):
    for cell in row:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")


def write_table(sheet, headers, rows):
    """Escribe encabezado + filas, da formato a fechas y dinero, y ajusta el ancho de columnas."""
    sheet.append(headers)
    style_header(sheet[1])
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"  # el encabezado queda fijo al bajar

    money_columns = {i for i, h in enumerate(headers) if "amount" in h.lower()}
    date_columns = {i for i, h in enumerate(headers) if "date" in h.lower()}
    for row in sheet.iter_rows(min_row=2):
        for i, cell in enumerate(row):
            if i in money_columns:
                cell.number_format = MONEY_FORMAT
            elif i in date_columns:
                cell.number_format = DATE_FORMAT
    autosize(sheet)


def autosize(sheet):
    """Ancho de cada columna según su texto más largo (con un tope)."""
    for column in sheet.columns:
        longest = 0
        for cell in column:
            if cell.value is None:
                continue
            if cell.number_format == MONEY_FORMAT:
                length = len(f"-${abs(cell.value):,.2f}")
            elif cell.number_format == DATE_FORMAT:
                length = 10
            else:
                length = len(str(cell.value))
            longest = max(longest, length)
        sheet.column_dimensions[column[0].column_letter].width = min(longest + 3, MAX_COLUMN_WIDTH)


def write_reconciliation_lines(sheet, recon, start_row):
    """Los renglones del formato estándar (saldo, + depósitos, − cheques, ...). Devuelve la fila siguiente."""
    row = start_row
    for label, amount, kind in recon.lines():
        if label in ("Book ending balance", "Difference"):
            row += 1   # renglón en blanco entre bloques
        sheet.cell(row=row, column=1, value=label)
        cell = sheet.cell(row=row, column=2, value=amount if amount is not None else "not given")
        if amount is not None:
            cell.number_format = MONEY_FORMAT
        if kind in ("total", "difference"):
            sheet.cell(row=row, column=1).font = Font(bold=True)
            cell.font = Font(bold=True)
        row += 1
    status = ("Enter both ending balances to get the difference" if not recon.is_complete
              else "RECONCILED" if recon.is_balanced else "NOT RECONCILED")
    sheet.cell(row=row, column=1, value=status).font = Font(bold=True)
    row += 1
    for hint in recon.hints():
        sheet.cell(row=row, column=1, value=hint)
        row += 1
    return row + 1


def write_summary(sheet, result, recon):
    sheet["A1"] = "Bank reconciliation"
    sheet["A1"].font = Font(bold=True, size=14)
    row = write_reconciliation_lines(sheet, recon, start_row=3)

    counts = [
        ("Matched", len(result.matched)),
        ("Needs review", len(result.review)),
        ("Bank only (not in the books)", len(result.bank_only)),
        ("Books only (not in the bank)", len(result.books_only)),
        ("Lines that could not be read", len(result.bad_lines)),
    ]
    for label, value in counts:
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=value)
        row += 1

    # Si hubo líneas que no se pudieron leer, se listan abajo con su motivo
    if result.bad_lines:
        row += 1
        sheet.cell(row=row, column=1, value="Lines that could not be read").font = Font(bold=True)
        for bad in result.bad_lines:
            row += 1
            file_name = "Bank" if bad.source == "bank" else "Books"
            sheet.cell(row=row, column=1, value=f"{file_name} file, line {bad.line}")
            sheet.cell(row=row, column=2, value=bad.reason)
            sheet.cell(row=row, column=3, value=bad.raw)
    autosize(sheet)
    sheet.column_dimensions["A"].width = 40


def write_reconciliation_sheet(sheet, recon):
    """El formato estándar con el detalle de cada partida, listo para imprimir.

    Columnas: A fecha, B referencia, C descripción, D monto de cada partida, E totales.
    """
    sheet["A1"] = "Bank reconciliation"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Prepared with Cuadrador. Items marked (pending review) are still open."
    sheet["A2"].font = Font(italic=True, color="4B5563")
    row = 4

    def label(text, amount=None, bold=False, border=None, font=None):
        """Un renglón con texto en A y un total en E."""
        nonlocal row
        sheet.cell(row=row, column=1, value=text).font = font or Font(bold=bold)
        if amount is not None or text.startswith(("=", "Difference")):
            cell = sheet.cell(row=row, column=5, value=amount if amount is not None else "not given")
            cell.number_format = MONEY_FORMAT
            cell.font = font or Font(bold=bold)
            if border:
                cell.border = border
        row += 1

    def items(title, lines, total):
        """El título de la sección, cada partida (columna D) y su total (columna E)."""
        nonlocal row
        label(title, total)
        for t in lines:
            sheet.cell(row=row, column=1, value=t.date).number_format = DATE_FORMAT
            sheet.cell(row=row, column=2, value=t.reference or (f"line {t.line}" if t.source == "bank" else ""))
            text = t.description + ("  (pending review)" if recon.is_pending(t) else "")
            sheet.cell(row=row, column=3, value=text)
            sheet.cell(row=row, column=4, value=t.amount).number_format = MONEY_FORMAT
            row += 1

    label("Bank ending balance", recon.bank_balance, bold=True)
    items("Add: deposits in transit", recon.deposits_in_transit, recon.deposits_total)
    items("Less: outstanding checks", recon.outstanding_checks, recon.checks_total)
    label("= Adjusted bank balance", recon.adjusted_bank, bold=True, border=THIN_TOP)
    row += 1
    label("Book ending balance", recon.book_balance, bold=True)
    items("Less: bank fees not recorded", recon.bank_fees, recon.fees_total)
    items("Add: interest not recorded", recon.interest, recon.interest_total)
    items("Add or less: other bank-only items", recon.other_bank, recon.other_total)
    label("= Adjusted book balance", recon.adjusted_book, bold=True, border=THIN_TOP)
    row += 1
    status_font = None if not recon.is_complete else GOOD_FONT if recon.is_balanced else BAD_FONT
    label("Difference", recon.difference, bold=True, border=DOUBLE_BOTTOM, font=status_font)
    status = ("Enter both ending balances to get the difference." if not recon.is_complete
              else "Reconciled: the difference is $0.00." if recon.is_balanced else "Not reconciled.")
    sheet.cell(row=row, column=1, value=status).font = status_font or Font(bold=True)
    row += 1
    for hint in recon.hints():
        sheet.cell(row=row, column=1, value=hint).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        sheet.row_dimensions[row].height = 30
        row += 1

    for column, width in zip("ABCDE", [30, 12, 44, 15, 16]):
        sheet.column_dimensions[column].width = width

    # Impresión: vertical, todo el ancho en una página, número de página abajo
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.5
    sheet.oddFooter.center.text = "Page &P of &N"
    sheet.print_area = f"A1:E{row}"


def build_workbook(result, recon):
    workbook = Workbook()
    write_reconciliation_sheet(workbook.active, recon)
    workbook.active.title = "Reconciliation"
    write_summary(workbook.create_sheet("Summary"), result, recon)

    write_table(workbook.create_sheet("Matched"), MATCH_HEADERS,
                [row for match in result.matched for row in match_rows(match)])
    write_table(workbook.create_sheet("Review"), MATCH_HEADERS,
                [row for match in result.review for row in match_rows(match)])
    write_table(workbook.create_sheet("Unmatched bank"),
                ["Bank line", "Date", "Description", "Amount", "Why"],
                [[t.line, t.date, t.description, t.amount, recon.explain(t)] for t in result.bank_only])
    write_table(workbook.create_sheet("Unmatched books"),
                ["Book line", "Date", "Name", "Ref", "Amount", "Why"],
                [[t.line, t.date, t.description, t.reference, t.amount, recon.explain(t)]
                 for t in result.books_only])
    return workbook


def save_report(result, recon, destination):
    """Guarda el Excel. `destination` puede ser una ruta o un archivo en memoria (BytesIO)."""
    build_workbook(result, recon).save(destination)
